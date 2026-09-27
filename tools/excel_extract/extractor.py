# -*- coding: utf-8 -*-
"""xlsx 数据提取核心逻辑。

流程：
    1. 选择 xlsx 文件
    2. 选择工作表
    3. 定位表头（支持多行表头，指定起止行；兼容合并单元格）
    4. 选择用于筛选的列
    5. 指定匹配方式（精确 / 包含 / 不等于 / 为空 / 非空）与目标值
    6. 预览命中数量，导出到新的 xlsx（表头原样保留，数据接在其后）

依赖 openpyxl（仅需读取/写入即可，不需要 pandas）。
"""

import os

from core import ui, paths

# ---------------------------------------------------------------- 依赖检查

def _load_openpyxl():
    """延迟导入 openpyxl，缺失时给出安装指引。"""
    try:
        import openpyxl
        return openpyxl
    except ImportError:
        ui.error("未检测到 openpyxl 库，无法处理 xlsx 文件。")
        ui.info("请先安装：pip install openpyxl")
        return None


# ---------------------------------------------------------------- 文件选择

def _list_dir_files(directory, exts=(".xlsx", ".xlsm")):
    """列出目录下指定后缀的文件。"""
    try:
        return sorted(
            f for f in os.listdir(directory)
            if f.lower().endswith(exts) and not f.startswith("~$")
        )
    except OSError:
        return []


def _choose_file():
    """让用户输入 xlsx 文件路径；Android 模式下可先浏览常见目录。

    返回绝对路径或 None（取消）。
    """
    ui.subtitle("步骤 1 / 6：导入 xlsx 文件")
    # 「浏览常见目录」只在 Android 模式提供：Windows 下不存在 /sdcard 等目录
    if paths.has_common_dirs():
        print("  直接输入完整路径，或先浏览常见目录。")
    else:
        print("  请输入 xlsx 文件的完整路径（路径含空格时可用引号包住）。")
    print(f"  示例：{paths.example_path()}")

    while True:
        if paths.has_common_dirs() and ui.ask_yes_no("是否浏览常见目录", "n"):
            dirs = paths.common_dirs()
            pick = ui.menu(
                [(str(i + 1), d) for i, d in enumerate(dirs)],
                title_text="选择目录",
                allow_back=True,
            )
            if pick is None:
                continue
            directory = dirs[int(pick) - 1]
            if not os.path.isdir(directory):
                ui.error(f"目录不存在：{directory}")
                continue
            files = _list_dir_files(directory)
            if not files:
                ui.warn(f"{directory} 中没有 xlsx 文件。")
                continue
            fpick = ui.menu(
                [(str(i + 1), f) for i, f in enumerate(files)],
                title_text=f"{directory} 中的表格",
                allow_back=True,
            )
            if fpick is None:
                continue
            return os.path.join(directory, files[int(fpick) - 1])

        raw = ui.ask("文件路径（直接回车取消）")
        p = paths.normalize_path(raw)
        if not p:
            ui.warn("已取消。")
            return None
        if not os.path.isfile(p):
            ui.error("文件不存在，请重新输入。")
            continue
        if not paths.is_readable_file(p):
            ui.error("文件不可读，请检查权限。")
            continue
        if not p.lower().endswith((".xlsx", ".xlsm")):
            ui.warn("该文件不是 .xlsx/.xlsm，可能无法解析。")
            if not ui.ask_yes_no("仍要继续吗", "n"):
                continue
        return p


# ---------------------------------------------------------------- 工作表选择

def _choose_sheet(wb):
    """选择工作表名。返回名称或 None（取消）。"""
    ui.subtitle("步骤 2 / 6：选择工作表")
    names = wb.sheetnames
    items = []
    for i, name in enumerate(names):
        ws = wb[name]
        items.append((str(i + 1), name, f"{ws.max_row} 行 × {ws.max_column} 列"))
    pick = ui.menu(items, title_text="工作表列表", allow_back=True)
    if pick is None:
        return None
    return names[int(pick) - 1]


# ---------------------------------------------------------------- 表头相关

def _col_letter(idx):
    """1 -> A, 27 -> AA。"""
    letters = ""
    while idx > 0:
        idx, rem = divmod(idx - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def _cell_text(value):
    """把单元格值转成用于显示/比较的字符串。"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _merged_value_map(ws):
    """收集合并单元格的「左上角值」，返回 {(row, col): 左上角值}。

    多行表头常见做法：同一列名用合并单元格纵向跨多行，
    或横向合并作为分组标题。这里把合并区域所有格都映射到左上角值，
    以便拼接表头时不会丢失文字。
    """
    mapping = {}
    for rng in ws.merged_cells.ranges:
        anchor = ws.cell(row=rng.min_row, column=rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                mapping[(r, c)] = anchor
    return mapping


def _header_cell_text(ws, merged, row, col):
    """读取表头单元格文字，兼容合并单元格。"""
    if (row, col) in merged:
        return _cell_text(merged[(row, col)])
    return _cell_text(ws.cell(row=row, column=col).value)


def _build_column_name(parts):
    """把多行表头片段拼接成列名。

    规则：过滤空片段；若相邻片段相同（常见于合并单元格向下填充）则去重。
    """
    cleaned = []
    for p in parts:
        p = (p or "").strip()
        if not p:
            continue
        if cleaned and cleaned[-1] == p:
            continue
        cleaned.append(p)
    return " / ".join(cleaned)


def _read_header(ws, header_rows):
    """读取表头，支持多行。

    header_rows: 表头行号列表（自上而下）。
    返回 [(列序号, 列名, 原始片段列表), ...]，跳过整段表头全空的列。
    """
    merged = _merged_value_map(ws)
    cols = []
    for c in range(1, ws.max_column + 1):
        parts = [_header_cell_text(ws, merged, r, c) for r in header_rows]
        name = _build_column_name(parts)
        if name:
            cols.append((c, name, parts))
    return cols


def _ask_header_rows(ws):
    """询问表头占用哪些行，支持多行表头。返回行号列表。"""
    ui.subtitle("步骤 3 / 6：定位表头")
    print("  单行表头：直接回车（默认第 1 行）。")
    print("  多行表头：输入起止行号，例如起 1 止 2，表示表头占第 1~2 行。")

    start = ui.ask_int("表头起始行", default=1, min_value=1, max_value=max(1, ws.max_row))
    if start is None:
        start = 1

    end = ui.ask_int(
        "表头结束行", default=start,
        min_value=start, max_value=max(1, ws.max_row),
    )
    if end is None:
        end = start

    return list(range(start, end + 1))


# ---------------------------------------------------------------- 条件构造

MATCH_MODES = [
    ("1", "精确匹配", "单元格内容与目标值完全相等，如 15 == 15"),
    ("2", "包含关键词", "单元格内容包含目标值，如 15班 含 15"),
    ("3", "不等于", "单元格内容与目标值不相等"),
    ("4", "为空", "单元格为空（无目标值）"),
    ("5", "非空", "单元格有内容（无目标值）"),
]


def _ask_match_mode():
    """询问匹配方式，返回 key。"""
    pick = ui.menu(MATCH_MODES, title_text="选择匹配方式", allow_back=True)
    if pick is None:
        return None
    return pick


def _ask_target_value(mode):
    """根据匹配方式询问目标值；为空/非空时返回 None。"""
    if mode in ("4", "5"):
        return None
    return ui.ask("目标值")


def _match(text, mode, target):
    """执行单值匹配判定。text/target 均为字符串。"""
    if mode == "1":
        return text == target
    if mode == "2":
        return target in text
    if mode == "3":
        return text != target
    if mode == "4":
        return text == ""
    if mode == "5":
        return text != ""
    return False


# ---------------------------------------------------------------- 主流程

def run():
    """工具入口。"""
    openpyxl = _load_openpyxl()
    if openpyxl is None:
        ui.pause()
        return

    # 1. 选文件
    file_path = _choose_file()
    if not file_path:
        return

    try:
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=False)
    except Exception as exc:  # noqa: BLE001
        ui.error(f"打开文件失败：{exc}")
        ui.pause()
        return

    # 2. 选工作表
    sheet_name = _choose_sheet(wb)
    if sheet_name is None:
        return
    ws = wb[sheet_name]

    # 3. 表头行（支持多行）
    header_rows = _ask_header_rows(ws)
    cols = _read_header(ws, header_rows)
    if not cols:
        ui.error("未在所选表头行读取到任何列名，请确认表头行范围是否正确。")
        ui.pause()
        return

    header_desc = (
        f"第 {header_rows[0]} 行"
        if len(header_rows) == 1
        else f"第 {header_rows[0]}~{header_rows[-1]} 行（共 {len(header_rows)} 行）"
    )
    ui.info(f"表头范围：{header_desc}，识别到 {len(cols)} 列。")
    print()
    print("  识别出的列名：")
    for i, (c, name, _) in enumerate(cols):
        print(f"    [{i + 1}] {name}  （{_col_letter(c)} 列）")

    # 4. 选筛选列
    ui.subtitle("步骤 4 / 6：选择筛选列")
    col_items = [
        (str(i + 1), name, f"第 {_col_letter(c)} 列")
        for i, (c, name, _) in enumerate(cols)
    ]
    pick = ui.menu(col_items, title_text="按哪一列筛选", allow_back=True)
    if pick is None:
        return
    col_index, col_name, _ = cols[int(pick) - 1]

    # 5. 匹配方式 + 目标值
    mode = _ask_match_mode()
    if mode is None:
        return
    target = _ask_target_value(mode)
    if target is None and mode not in ("4", "5"):
        ui.warn("目标值为空，已取消。")
        return
    target = _cell_text(target) if target is not None else None

    mode_label = dict((k, v) for k, v, _ in MATCH_MODES).get(mode, mode)

    # 开始扫描（数据从表头最后一行的下一行开始）
    data_start = header_rows[-1] + 1
    matched_rows = []
    for r in range(data_start, ws.max_row + 1):
        raw = ws.cell(row=r, column=col_index).value
        text = _cell_text(raw)
        if _match(text, mode, target):
            matched_rows.append(r)

    ui.subtitle("步骤 5 / 6：预览结果")
    ui.info(f"文件：{file_path}")
    ui.info(f"工作表：{sheet_name}")
    ui.info(f"筛选条件：{col_name}（{_col_letter(col_index)} 列） {mode_label} "
            f"{'「' + target + '」' if target is not None else ''}")
    # 表头行显示（多行表头时列出）
    if len(header_rows) > 1:
        ui.info(f"表头共 {len(header_rows)} 行：{header_rows}")
    ui.info(f"数据范围：第 {data_start} ~ {ws.max_row} 行，共 {len(matched_rows)} 行命中。")

    if not matched_rows:
        ui.warn("没有任何数据符合条件，未生成文件。")
        ui.pause()
        return

    preview_n = min(10, len(matched_rows))
    print()
    print(f"  命中行号预览（前 {preview_n} 条）：")
    print("  " + ", ".join(str(r) for r in matched_rows[:preview_n])
          + (" ..." if len(matched_rows) > preview_n else ""))

    # 6. 导出
    ui.subtitle("步骤 6 / 6：导出新文件")
    default_name = paths.default_output_name(file_path, f"{col_name}_{target or mode_label}")
    out_path = ui.ask("输出文件路径", default=os.path.join(paths.ensure_output_dir(), default_name))
    out_path = paths.normalize_path(out_path)
    if not out_path:
        ui.warn("输出路径为空，已取消。")
        return
    if not out_path.lower().endswith((".xlsx", ".xlsm")):
        out_path += ".xlsx"
    out_dir = os.path.dirname(os.path.abspath(out_path))
    try:
        os.makedirs(out_dir, exist_ok=True)
    except OSError as exc:
        ui.error(f"无法创建输出目录：{exc}")
        ui.pause()
        return

    try:
        out_wb = openpyxl.Workbook()
        out_ws = out_wb.active
        out_ws.title = sheet_name[:31] or "结果"

        header_count = len(header_rows)

        # 1) 原样复制所有表头行（含合并单元格内容）
        src_merged = _merged_value_map(ws)
        for out_r, src_r in enumerate(header_rows, start=1):
            for c in range(1, ws.max_column + 1):
                out_ws.cell(row=out_r, column=c).value = _header_cell_text(
                    ws, src_merged, src_r, c
                )

        # 2) 命中行按序复制到表头之后
        for i, r in enumerate(matched_rows):
            for c in range(1, ws.max_column + 1):
                out_ws.cell(row=header_count + 1 + i, column=c).value = \
                    ws.cell(row=r, column=c).value

        out_wb.save(out_path)
    except PermissionError:
        ui.error("写入被拒绝，目标文件可能正被占用，或路径无写权限。")
        ui.pause()
        return
    except Exception as exc:  # noqa: BLE001
        ui.error(f"导出失败：{exc}")
        ui.pause()
        return

    print()
    ui.success(f"已生成：{out_path}")
    header_note = f"{len(header_rows)} 行表头" if len(header_rows) > 1 else "1 行表头"
    ui.info(f"共写入 {len(matched_rows)} 行数据 + {header_note}。")

    if ui.ask_yes_no("是否继续在本次结果上再次筛选", "n"):
        _run_on_new_file(out_path)

    ui.pause()


def _run_on_new_file(path):
    """对刚生成的文件递归调用一次筛选（简化：重新走一遍主流程）。"""
    print()
    ui.info(f"请在上方「文件路径」处输入：{path}")
    ui.pause("按回车继续...")
    run()
